from __future__ import annotations

import io
from pathlib import Path

import pymupdf
import pytesseract
from docx import Document
from PIL import Image


class UnsupportedFileTypeError(ValueError):
    pass


def extract_document(data: bytes, filename: str) -> dict:
    extension = Path(filename).suffix.lower().lstrip(".")

    if extension == "pdf":
        return _extract_pdf(data)
    if extension == "docx":
        return _extract_docx(data)
    # Resume imports intentionally accept documents only. Images and arbitrary
    # text files are rejected so upload behavior matches the product contract.
    if extension == "doc":
        raise UnsupportedFileTypeError(
            "Legacy .doc files are not supported. Save the document as .docx or PDF and upload it again."
        )
    raise UnsupportedFileTypeError(
        f"Unsupported file type: .{extension}. Upload a PDF or DOCX resume."
    )


def extract_text(data: bytes, filename: str) -> str:
    return extract_document(data, filename)["text"]


def _extract_pdf(data: bytes) -> dict:
    pages = []
    blocks = []

    with pymupdf.open(stream=data, filetype="pdf") as document:
        for page_number, page in enumerate(document, start=1):
            page_blocks = []
            for index, block in enumerate(page.get_text("blocks", sort=True)):
                text = str(block[4]).strip()
                if not text:
                    continue
                block_id = f"p{page_number}b{index + 1}"
                page_blocks.append(text)
                blocks.append(
                    {
                        "id": block_id,
                        "page": page_number,
                        "text": text,
                        "bbox": list(block[:4]),
                    }
                )

            # A mixed PDF may contain selectable text on some pages and scans
            # on others. OCR each text-empty page instead of only OCRing when
            # the entire document has no selectable text.
            if not page_blocks and page.get_images(full=True):
                try:
                    pixmap = page.get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False)
                    image = Image.open(io.BytesIO(pixmap.tobytes("png")))
                    ocr_text = pytesseract.image_to_string(image).strip()
                except pytesseract.TesseractNotFoundError as exc:
                    raise RuntimeError(
                        "Tesseract OCR is required because the PDF contains scanned pages."
                    ) from exc
                if ocr_text:
                    page_blocks.append(ocr_text)
                    blocks.append({
                        "id": f"ocr-p{page_number}",
                        "page": page_number,
                        "text": ocr_text,
                        "type": "ocr",
                    })

            links = []
            for link in page.get_links():
                uri = link.get("uri")
                if uri and uri.lower().startswith(("http://", "https://")):
                    links.append(uri)

            if links:
                # The URLs themselves must be part of the text; previously only a
                # "[Embedded links]" marker was added, so hyperlinked LinkedIn /
                # GitHub / project URLs never reached the extractors.
                page_blocks.append("[Embedded links]")
                for link_index, uri in enumerate(dict.fromkeys(links), start=1):
                    page_blocks.append(uri)
                    blocks.append(
                        {
                            "id": f"p{page_number}link{link_index}",
                            "page": page_number,
                            "text": uri,
                            "type": "link",
                        }
                    )

            page_text = "\n".join(page_blocks).strip()
            if page_text:
                pages.append(f"[PAGE {page_number}]\n{page_text}")

        text = "\n\n".join(pages).strip()

    if not text:
        raise ValueError("The document contains no readable text.")

    return {"text": text, "blocks": blocks}


def _extract_docx(data: bytes) -> dict:
    document = Document(io.BytesIO(data))
    blocks = []

    for section in document.sections:
        for paragraph in section.header.paragraphs:
            value = paragraph.text.strip()
            if value:
                blocks.append({"id": f"h{len(blocks)+1}", "page": 1, "text": value, "type": "header"})
        for paragraph in section.footer.paragraphs:
            value = paragraph.text.strip()
            if value:
                blocks.append({"id": f"f{len(blocks)+1}", "page": 1, "text": value, "type": "footer"})

    for paragraph in document.paragraphs:
        value = paragraph.text.strip()
        if value:
            blocks.append({"id": f"b{len(blocks)+1}", "page": 1, "text": value, "type": "paragraph"})

    for table_index, table in enumerate(document.tables, start=1):
        for row_index, row in enumerate(table.rows, start=1):
            values = [cell.text.strip() for cell in row.cells]
            if any(values):
                blocks.append(
                    {
                        "id": f"t{table_index}r{row_index}",
                        "page": 1,
                        "text": " | ".join(values),
                        "type": "table",
                    }
                )

    seen_links: set[str] = set()
    for rel in document.part.rels.values():
        target = getattr(rel, "target_ref", "") if getattr(rel, "is_external", False) else ""
        if (
            rel.reltype.endswith("/hyperlink")
            and target.lower().startswith(("http://", "https://"))
            and target not in seen_links
        ):
            seen_links.add(target)
            blocks.append({"id": f"l{len(seen_links)}", "page": 1, "text": target, "type": "link"})

    text = "\n".join(block["text"] for block in blocks).strip()
    if not text:
        raise ValueError("The DOCX file contains no extractable text.")

    return {"text": text, "blocks": blocks}
